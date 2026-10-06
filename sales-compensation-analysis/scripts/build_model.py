"""Build the June 2026 commission model as a formula-driven Excel workbook.

Every number in the calculation and output tabs is a live formula that reads the input
tabs (copied unchanged from the case workbook) and the Control / Treatments tabs.
Run:  python scripts/build_model.py   then recalculate with LibreOffice (see README).
"""
from pathlib import Path
import datetime as dt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "Sales_Compensation_Analyst_Case_Study_ORIGINAL.xlsx"
OUT = ROOT / "deliverables" / "June_2026_Commission_Model.xlsx"

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SUB = Font(name=FONT, size=11, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="555555")
F_BOLD = Font(name=FONT, size=10, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_INPUT_HEAD = PatternFill("solid", fgColor="7F6000")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_ALT = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

USD = '$#,##0.00;($#,##0.00);"-"'
USD0 = '$#,##0;($#,##0);"-"'
NUM = '#,##0.00;(#,##0.00);"-"'
INT = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
DATE = "yyyy-mm-dd"
MULT = '0.00"x"'

INPUT_TABS = ["1. Sales Roster", "2. HR Export", "3. Exchange Rates", "4. Ramp Terms",
              "5. Compensation", "6. Accelerators", "7. Opportunity to ID Mapping",
              "8. Opportunity Export", "9. App Export", "10. Billing Export", "11. Slack Scrape"]

# Upper bounds for input ranges: generous so a larger July extract still fits.
MAX = {"9. App Export": 250000, "10. Billing Export": 20000, "8. Opportunity Export": 10000,
       "7. Opportunity to ID Mapping": 10000, "2. HR Export": 10000, "1. Sales Roster": 1000,
       "3. Exchange Rates": 1000}

N_REPS, N_ACCTS, N_OPPS = 80, 761, 815


def q(sheet):
    return "'" + sheet + "'"


def rng(sheet, col, maxrow=None):
    maxrow = maxrow or MAX[sheet]
    return f"{q(sheet)}!${col}$2:${col}${maxrow}"


NAMES = {
    # App Export (tab 9)
    "SH_ACCT": rng("9. App Export", "A"), "SH_ID": rng("9. App Export", "B"),
    "SH_PDATE": rng("9. App Export", "C"), "SH_WORKED": rng("9. App Export", "D"),
    "SH_WDATE": rng("9. App Export", "E"), "SH_FIRST": rng("9. App Export", "F"),
    # Billing Export (tab 10)
    "BL_ACCT": rng("10. Billing Export", "A"), "BL_MONTH": rng("10. Billing Export", "B"),
    "BL_TYPE": rng("10. Billing Export", "C"), "BL_SHIFT": rng("10. Billing Export", "D"),
    "BL_AMT": rng("10. Billing Export", "E"), "BL_STATUS": rng("10. Billing Export", "F"),
    # Opportunity Export (tab 8)
    "OP_ACCT": rng("8. Opportunity Export", "A"), "OP_ID": rng("8. Opportunity Export", "B"),
    "OP_CLOSER": rng("8. Opportunity Export", "C"), "OP_OWNER": rng("8. Opportunity Export", "D"),
    "OP_TYPE": rng("8. Opportunity Export", "E"), "OP_STAGE": rng("8. Opportunity Export", "F"),
    "OP_CWDATE": rng("8. Opportunity Export", "G"), "OP_EFF": rng("8. Opportunity Export", "H"),
    "OP_SEG": rng("8. Opportunity Export", "I"), "OP_CHARGE": rng("8. Opportunity Export", "J"),
    "OP_PAY": rng("8. Opportunity Export", "K"), "OP_TERMS": rng("8. Opportunity Export", "L"),
    # HR Export (tab 2)
    "HR_ID": rng("2. HR Export", "A"), "HR_ROLE": rng("2. HR Export", "B"),
    "HR_START": rng("2. HR Export", "C"), "HR_END": rng("2. HR Export", "D"),
    "HR_RSTART": rng("2. HR Export", "E"), "HR_REND": rng("2. HR Export", "F"),
    "HR_PREV": rng("2. HR Export", "G"), "HR_PSTART": rng("2. HR Export", "H"),
    "HR_PEND": rng("2. HR Export", "I"), "HR_MGR": rng("2. HR Export", "J"),
    "HR_CC": rng("2. HR Export", "K"),
    # Exchange rates (tab 3)
    "FX_CC": rng("3. Exchange Rates", "A"), "FX_CODE": rng("3. Exchange Rates", "D"),
    "FX_RATE": rng("3. Exchange Rates", "E"),
    # Compensation (tab 5)
    "CP_ROLE": "'5. Compensation'!$A$2:$A$100", "CP_METRIC1": "'5. Compensation'!$C$2:$C$100",
    "CP_QUOTA": "'5. Compensation'!$D$2:$D$100", "CP_REPORTS": "'5. Compensation'!$F$2:$F$100",
    "CP_RATE": "'5. Compensation'!$G$2:$G$100", "CP_METRIC2": "'5. Compensation'!$J$2:$J$100",
    "CP_QUOTA2": "'5. Compensation'!$K$2:$K$100", "CP_RATE2": "'5. Compensation'!$L$2:$L$100",
}


class Table:
    """Column-keyed formula table: templates reference own columns as {key} and row as {r}."""

    def __init__(self, ws, header_row, cols):
        self.ws, self.hr, self.cols = ws, header_row, cols
        self.L = {c[0]: get_column_letter(i + 1) for i, c in enumerate(cols)}

    def col(self, key):
        return self.L[key]

    def write(self, n_rows, first_row=None):
        ws, hr = self.ws, self.hr
        first_row = first_row or hr + 1
        for i, (key, header, tmpl, fmt, *rest) in enumerate(self.cols):
            c = ws.cell(row=hr, column=i + 1, value=header)
            c.font, c.fill = F_HEAD, FILL_HEAD
            c.alignment = Alignment(wrap_text=True, vertical="center")
            if rest and rest[0]:
                c.comment = Comment(rest[0], "model")
            ws.column_dimensions[get_column_letter(i + 1)].width = max(11, min(32, len(header) * 0.9))
        for k in range(n_rows):
            r = first_row + k
            for i, (key, header, tmpl, fmt, *rest) in enumerate(self.cols):
                v = tmpl.format(r=r, k=k, **self.L) if isinstance(tmpl, str) else tmpl
                cell = ws.cell(row=r, column=i + 1, value=v)
                cell.font = F_BASE
                if fmt:
                    cell.number_format = fmt
        ws.freeze_panes = ws.cell(row=hr + 1, column=3)
        ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(self.cols))}{first_row + n_rows - 1}"
        return first_row, first_row + n_rows - 1


def copy_inputs(wb_out):
    src = openpyxl.load_workbook(SRC)
    for name in INPUT_TABS:
        s = src[name]
        t = wb_out.create_sheet(name)
        for row in s.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                n = t.cell(row=c.row, column=c.column, value=c.value)
                if c.row == 1 or (name == "6. Accelerators" and isinstance(c.value, str)):
                    n.font = Font(name=FONT, size=10, bold=(c.row in (1, 2, 7, 8) or name != "6. Accelerators"),
                                  color="FFFFFF" if c.row == 1 and name != "6. Accelerators" else "0000FF")
                    if c.row == 1 and name != "6. Accelerators":
                        n.fill = FILL_INPUT_HEAD
                if isinstance(c.value, (dt.datetime, dt.date)):
                    n.number_format = DATE
        for col in range(1, s.max_column + 1):
            t.column_dimensions[get_column_letter(col)].width = 20
        t.freeze_panes = "A2"
        t.sheet_properties.tabColor = "7F6000"
    return src


def build():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    order = ["README", "Summary", "Exceptions_Ledger", "Scenario_Accelerator", "Plan_Observations",
             "Checks", "Control", "Treatments", "Calc_People", "Calc_Opps", "Calc_Accounts",
             "Pay_AsWritten", "Pay_Recommended"]
    S = {n: wb.create_sheet(n) for n in order}
    src = copy_inputs(wb)

    # locate rows on tab 6 (accelerators / SPIFF) instead of assuming positions
    t6 = src["6. Accelerators"]
    acc_rows, spiff_rows = {}, {}
    for row in t6.iter_rows(min_row=1):
        v = row[0].value
        if v in ("SMB Account Executive", "ENT Account Executive"):
            acc_rows[v] = row[0].row
        if v in ("SMB Account Manager", "ENT Account Manager"):
            spiff_rows[v] = row[0].row
    a0, a1 = min(acc_rows.values()), max(acc_rows.values())
    s0, s1 = min(spiff_rows.values()), max(spiff_rows.values())
    NAMES.update({
        "AC_ROLE": f"'6. Accelerators'!$A${a0}:$A${a1}", "AC_T1": f"'6. Accelerators'!$D${a0}:$D${a1}",
        "AC_M1": f"'6. Accelerators'!$F${a0}:$F${a1}", "AC_T2": f"'6. Accelerators'!$H${a0}:$H${a1}",
        "AC_M2": f"'6. Accelerators'!$I${a0}:$I${a1}",
        "SP_ROLE": f"'6. Accelerators'!$A${s0}:$A${s1}", "SP_AMT": f"'6. Accelerators'!$D${s0}:$D${s1}",
    })

    # ------------------------------------------------------------------ Control
    ws = S["Control"]
    ws.sheet_properties.tabColor = "FFC000"
    ws["A1"], ws["A1"].font = "Control — run settings and plan parameters", F_TITLE
    ws["A2"] = "Blue cells are inputs. Everything else in the model reads from here or from the input tabs. To run another month, change B5 only and paste the new extracts into tabs 1–11."
    ws["A2"].font = F_NOTE
    rows = [
        # (row, label, value/formula, fmt, name, note, is_input)
        (5, "Comp month (first day)", dt.datetime(2026, 6, 1), DATE, "CM_START", "The month being paid. Change this to rerun.", True),
        (6, "Comp month end", "=EOMONTH(CM_START,0)", DATE, "CM_END", None, False),
        (7, "Days in comp month", "=CM_END-CM_START+1", INT, "CM_DAYS", "Quota and ramp prorate by calendar day (7.10, 7.11).", False),
        (8, "Prior month start", "=EDATE(CM_START,-1)", DATE, "PM_START", "Used by the retention SPIFF (7.12).", False),
        (9, "Prior month end", "=CM_START-1", DATE, "PM_END", None, False),
        (10, "Comp month label (billing_month format)", '=YEAR(CM_START)&"-"&TEXT(MONTH(CM_START),"00")', None, "CM_LABEL", "Matches tab 10 billing_month (7.2).", False),
        (11, "Prior month label", '=YEAR(PM_START)&"-"&TEXT(MONTH(PM_START),"00")', None, "PM_LABEL", None, False),
        (14, "AE window length (days)", 30, INT, "WIN_LEN", "7.4: three consecutive 30-day windows from the first worked shift.", True),
        (15, "Credit factor for F31-60D and F61-90D", 0.5, PCT, "LATER_FACTOR", "7.4: later windows pay at half, applied to the credit.", True),
        (16, "GTM payout rate (on GSV)", 0.03, PCT, "GTM_RATE", "7.6: GTM dollars pay 3%; converted to credited revenue as GSV x 3% / role rate.", True),
        (17, "Ramp month length (days)", 30, INT, "RAMP_LEN", "7.10: month 1 = days 1-30, month 2 = days 31-60.", True),
        (18, "Rounding: decimals for payout", 2, INT, "DEC", "A12: full precision throughout; round USD once per rep, convert, round local.", True),
        (19, "Approved ticket status value", "approved", None, "ST_APPROVED", "7.2: only approved tickets reduce Existing revenue.", True),
        (20, "Closed Won stage value", "Closed Won", None, "ST_WON", "7.3: opportunity_stage governs.", True),
        (21, "GTM account_type value", "GTM", None, "GTM_TYPE", "A17: account_type = GTM is the result of the Applied Science tag.", True),
    ]
    for r, label, val, fmt, name, note, is_in in rows:
        ws.cell(row=r, column=1, value=label).font = F_BASE
        c = ws.cell(row=r, column=2, value=val)
        c.font = F_INPUT if is_in else F_BASE
        if is_in and r == 5:
            c.fill = FILL_KEY
        if fmt:
            c.number_format = fmt
        if note:
            ws.cell(row=r, column=3, value=note).font = F_NOTE
        wb.defined_names[name] = DefinedName(name, attr_text=f"Control!$B${r}")
    ws["A4"], ws["A4"].font = "Run settings", F_SUB
    ws["A13"], ws["A13"].font = "Plan parameters transcribed from Section 07 (cite the clause)", F_SUB

    ws["E4"], ws["E4"].font = "Net terms multipliers (7.5)", F_SUB
    for i, (t, m) in enumerate([(15, 1.2), (30, 1.0), (45, 0.9), (60, 0.7)]):
        ws.cell(row=5 + i, column=5, value=t).font = F_INPUT
        c = ws.cell(row=5 + i, column=6, value=m); c.font = F_INPUT; c.number_format = PCT
    ws["E9"] = "Net terms on a new-logo account scale its revenue (never the rate, never Existing, never GTM)."
    ws["E9"].font = F_NOTE
    wb.defined_names["TERMS_DAYS"] = DefinedName("TERMS_DAYS", attr_text="Control!$E$5:$E$8")
    wb.defined_names["TERMS_MULT"] = DefinedName("TERMS_MULT", attr_text="Control!$F$5:$F$8")

    ws["E12"], ws["E12"].font = "Ramp credit % by ramp month (tab 4 text)", F_SUB
    ws["E13"], ws["F13"], ws["G13"] = "Role", "Month 1", "Month 2"
    for c in ("E13", "F13", "G13"):
        ws[c].font = F_BOLD
    for i, (role, m1, m2) in enumerate([("SMB Account Executive", 1, 0.5), ("ENT Account Executive", 1, 0.5),
                                        ("SMB Account Manager", 1, 1), ("ENT Account Manager", 1, 1)]):
        ws.cell(row=14 + i, column=5, value=role).font = F_INPUT
        for j, v in enumerate((m1, m2)):
            c = ws.cell(row=14 + i, column=6 + j, value=v); c.font = F_INPUT; c.number_format = PCT
    ws["E18"] = "Source: tab 4. Ramp Terms, column D ('100% quota credit in days 1-30 and 50%/100% in days 31-60'). Roles not listed do not ramp."
    ws["E18"].font = F_NOTE
    for nm, col in (("RAMP_ROLE", "E"), ("RAMP_M1", "F"), ("RAMP_M2", "G")):
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"Control!${col}$14:${col}$17")

    ws["E21"], ws["E21"].font = "Scenario lever — SMB AE accelerator (task 3.3)", F_SUB
    ws["E22"], ws["F22"], ws["G22"] = "", "Band I", "Band II"
    ws["E23"] = "Current (linked to tab 6)"
    ws["F23"] = f"=INDEX(AC_M1,MATCH(\"SMB Account Executive\",AC_ROLE,0))"
    ws["G23"] = f"=INDEX(AC_M2,MATCH(\"SMB Account Executive\",AC_ROLE,0))"
    ws["E24"] = "Proposed (VP of Sales)"
    ws["F24"], ws["G24"] = 1.3, 1.5
    for c in ("F23", "G23", "F24", "G24"):
        ws[c].number_format = MULT
        ws[c].font = F_INPUT if c in ("F24", "G24") else F_BASE
    ws["F24"].fill = ws["G24"].fill = FILL_KEY
    wb.defined_names["SCN_M1"] = DefinedName("SCN_M1", attr_text="Control!$F$24")
    wb.defined_names["SCN_M2"] = DefinedName("SCN_M2", attr_text="Control!$G$24")
    ws.column_dimensions["A"].width = 40; ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 60; ws.column_dimensions["E"].width = 26
    for col in "FG":
        ws.column_dimensions[col].width = 12

    for n, ref in NAMES.items():
        wb.defined_names[n] = DefinedName(n, attr_text=ref)

    # ------------------------------------------------------------------ Treatments
    ws = S["Treatments"]
    ws.sheet_properties.tabColor = "FFC000"
    ws["A1"], ws["A1"].font = "Treatments — sign-off decisions applied in the Recommended layer only", F_TITLE
    ws["A2"] = ("Data errors are corrected by plan rules in both layers. This tab holds payouts the plan as written "
                "would make but that we would not sign off (A1, A18). Set Active = 0 to switch a treatment off; "
                "Pay_Recommended then equals Pay_AsWritten for that account.")
    ws["A2"].font = F_NOTE
    hdr = ["crm_account_id", "Treatment code", "Active (1/0)", "Ledger ref", "What it does", "Why", "Decision owner"]
    for i, h in enumerate(hdr):
        c = ws.cell(row=4, column=i + 1, value=h); c.font, c.fill = F_HEAD, FILL_HEAD
    trows = [("b2b39ae1-3d7a-47ee-962b-f400aabcd70c", "GTM_WORKED_ONLY", 1, "E01",
              "Credits GTM GSV on worked shifts only; unworked postings are held, not forfeited.",
              "313 shifts posted (293 in June), 45 worked (14% fill vs ~78% norm). Slack: AE loading shift templates for rural sites; facility DON says they post a handful a day; ops flagged an open template. Approved as A18.",
              "VP of Sales / Head of Sales & AM")]
    for i, tr in enumerate(trows):
        for j, v in enumerate(tr):
            c = ws.cell(row=5 + i, column=j + 1, value=v)
            c.font = F_INPUT
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFG", (40, 20, 12, 10, 45, 70, 24)):
        ws.column_dimensions[col].width = w
    wb.defined_names["TR_ACCT"] = DefinedName("TR_ACCT", attr_text="Treatments!$A$5:$A$100")
    wb.defined_names["TR_CODE"] = DefinedName("TR_CODE", attr_text="Treatments!$B$5:$B$100")
    wb.defined_names["TR_ACTIVE"] = DefinedName("TR_ACTIVE", attr_text="Treatments!$C$5:$C$100")

    # ------------------------------------------------------------------ Calc_People
    ws = S["Calc_People"]
    ws.sheet_properties.tabColor = "548235"
    ws["A1"], ws["A1"].font = "Calc_People — one row per rep on tab 1; HR governs role, dates and reporting line (A16)", F_SUB
    RO = q("1. Sales Roster")
    HRM = "MATCH($C{r},HR_ID,0)"
    pcols = [
        ("src", "Roster row", "=ROW(" + RO + "!A{r})", INT),
        ("name", "Name", "=" + RO + "!A{r}", None),
        ("hr", "hr_id", "=" + RO + "!B{r}", None),
        ("crm", "crm_owner_id", "=" + RO + "!C{r}", None),
        ("slack", "slack_user_id", "=" + RO + "!D{r}", None),
        ("hrrow", "HR row", "=IFERROR(" + HRM.replace("$C{r}", "{hr}{r}") + "+1,\"MISSING\")", INT),
        ("role", "Current Role (HR)", "=INDEX(HR_ROLE,{hrrow}{r}-1)", None),
        ("es", "Employment start", "=INDEX(HR_START,{hrrow}{r}-1)", DATE),
        ("ee", "Employment end", "=IF(INDEX(HR_END,{hrrow}{r}-1)=\"\",\"\",INDEX(HR_END,{hrrow}{r}-1))", DATE),
        ("rs", "Current Role Start", "=INDEX(HR_RSTART,{hrrow}{r}-1)", DATE),
        ("pe", "Previous Role End", "=IF(INDEX(HR_PEND,{hrrow}{r}-1)=\"\",\"\",INDEX(HR_PEND,{hrrow}{r}-1))", DATE),
        ("rdflag", "Role-date overlap (A19)", "=IF({pe}{r}=\"\",\"\",IF({pe}{r}>={rs}{r},\"CHECK\",\"\"))", None),
        ("mgr", "Manager - ID (L1)", "=INDEX(HR_MGR,{hrrow}{r}-1)", None),
        ("mgrname", "Manager name", "=IFERROR(INDEX($B$4:$B$83,MATCH({mgr}{r},$C$4:$C$83,0)),\"(outside export)\")", None),
        ("l2", "Chain L2", "=IFERROR(INDEX(HR_MGR,MATCH({mgr}{r},HR_ID,0)),\"\")", None),
        ("l3", "Chain L3", "=IFERROR(INDEX(HR_MGR,MATCH({l2}{r},HR_ID,0)),\"\")", None),
        ("l4", "Chain L4", "=IFERROR(INDEX(HR_MGR,MATCH({l3}{r},HR_ID,0)),\"\")", None),
        ("cc", "Country Code", "=INDEX(HR_CC,{hrrow}{r}-1)", None),
        ("cur", "Currency", "=IFERROR(INDEX(FX_CODE,MATCH({cc}{r},FX_CC,0)),\"NO RATE\")", None),
        ("fx", "FX: local per 1 USD", "=IFERROR(INDEX(FX_RATE,MATCH({cc}{r},FX_CC,0)),\"NO RATE\")", "0.0000"),
        ("cprow", "Comp row", "=MATCH({role}{r},CP_ROLE,0)", INT),
        ("m1", "Metric I", "=INDEX(CP_METRIC1,{cprow}{r})", None),
        ("m2", "Metric II", "=INDEX(CP_METRIC2,{cprow}{r})", None),
        ("isic", "Individual contributor?", "=INDEX(CP_REPORTS,{cprow}{r})=\"na\"", None),
        ("isae", "IC on New (AE)?", "=AND({isic}{r},{m1}{r}=\"New Customer Net Revenue\")", None),
        ("isam", "IC on Existing (AM)?", "=AND({isic}{r},{m1}{r}=\"Existing Customer Net Revenue\")", None),
        ("quota", "Monthly quota (Metric I)", "=INDEX(CP_QUOTA,{cprow}{r})", USD0),
        ("quota2", "Monthly quota (Metric II)", "=IF(ISNUMBER(INDEX(CP_QUOTA2,{cprow}{r})),INDEX(CP_QUOTA2,{cprow}{r}),0)", USD0),
        ("rate", "Standard Payout (Metric I)", "=INDEX(CP_RATE,{cprow}{r})", "0.0000%"),
        ("rate2", "Standard Payout 2 (Metric II)", "=IF(AND({m2}{r}=\"Existing Customer Net Revenue\",ISNUMBER(INDEX(CP_RATE2,{cprow}{r}))),INDEX(CP_RATE2,{cprow}{r}),0)", "0.0000%"),
        ("eeff", "Employed through (in month)", "=IF({ee}{r}=\"\",CM_END,MIN({ee}{r},CM_END))", DATE),
        ("days", "Days employed in month", "=MAX(0,{eeff}{r}-MAX(CM_START,{es}{r})+1)", INT),
        ("qpr", "Quota prorated (7.11)", "={quota}{r}*{days}{r}/CM_DAYS", USD),
        ("qpr2", "Quota II prorated", "={quota2}{r}*{days}{r}/CM_DAYS", USD),
        ("p1", "Ramp % month 1", "=IFERROR(INDEX(RAMP_M1,MATCH({role}{r},RAMP_ROLE,0)),0)", PCT),
        ("p2", "Ramp % month 2", "=IFERROR(INDEX(RAMP_M2,MATCH({role}{r},RAMP_ROLE,0)),0)", PCT),
        ("d1", "Days in ramp month 1", "=IF({p1}{r}+{p2}{r}=0,0,MAX(0,MIN({eeff}{r},{rs}{r}+RAMP_LEN-1)-MAX(CM_START,{es}{r},{rs}{r})+1))", INT),
        ("d2", "Days in ramp month 2", "=IF({p1}{r}+{p2}{r}=0,0,MAX(0,MIN({eeff}{r},{rs}{r}+2*RAMP_LEN-1)-MAX(CM_START,{es}{r},{rs}{r}+RAMP_LEN)+1))", INT),
        ("ramping", "Ramping in month? (A10: any day)", "=({d1}{r}+{d2}{r})>0", None),
        ("rampcr", "Ramp credit floor (7.10)", "={quota}{r}/CM_DAYS*({d1}{r}*{p1}{r}+{d2}{r}*{p2}{r})", USD),
        ("t1", "Band I starts at", "=IFERROR(INDEX(AC_T1,MATCH({role}{r},AC_ROLE,0)),\"\")", USD0),
        ("t2", "Band II starts at", "=IFERROR(INDEX(AC_T2,MATCH({role}{r},AC_ROLE,0)),\"\")", USD0),
        ("am1", "Multiplier I", "=IFERROR(INDEX(AC_M1,MATCH({role}{r},AC_ROLE,0)),\"\")", MULT),
        ("am2", "Multiplier II", "=IFERROR(INDEX(AC_M2,MATCH({role}{r},AC_ROLE,0)),\"\")", MULT),
        ("spamt", "Retention SPIFF amount", "=IFERROR(INDEX(SP_AMT,MATCH({role}{r},SP_ROLE,0)),0)", USD0),
    ]
    pt = Table(ws, 3, [(k, h, t.replace("{r}", "{r}"), f) for k, h, t, f in pcols])
    # fix the roster row reference: roster data row = sheet row - 2 (header row 3 -> roster row 2)
    for i, c in enumerate(pt.cols):
        key, h, t, f = c
        pt.cols[i] = (key, h, t.replace(RO + "!A{r}", RO + "!A{rr}").replace(RO + "!B{r}", RO + "!B{rr}")
                      .replace(RO + "!C{r}", RO + "!C{rr}").replace(RO + "!D{r}", RO + "!D{rr}"), f)
    _write_people(pt, N_REPS)
    P = pt.L
    PR0, PR1 = 4, 3 + N_REPS
    for key in ("name", "hr", "crm", "role", "rate", "es", "eeff", "days", "l1", "l2", "l3", "l4", "isae", "isam",
                "ramping", "rampcr", "qpr", "qpr2", "quota", "rate2", "m1", "m2", "t1", "t2", "am1", "am2",
                "spamt", "fx", "cur", "mgr", "isic", "ee", "mgrname", "cc"):
        k2 = "mgr" if key == "l1" else key
        wb.defined_names["PP_" + key.upper()] = DefinedName(
            "PP_" + key.upper(), attr_text=f"Calc_People!${P[k2]}${PR0}:${P[k2]}${PR1}")

    # ------------------------------------------------------------------ Calc_Opps
    ws = S["Calc_Opps"]
    ws.sheet_properties.tabColor = "548235"
    ws["A1"], ws["A1"].font = "Calc_Opps — one row per opportunity on tab 8; finds each account's go-live (7.3)", F_SUB
    OP = q("8. Opportunity Export")
    ocols = [
        ("src", "Tab 8 row", "=ROW(" + OP + "!A{rr})", INT),
        ("acct", "crm_account_id", "=" + OP + "!A{rr}", None),
        ("opp", "opportunity_id", "=" + OP + "!B{rr}", None),
        ("stage", "opportunity_stage", "=" + OP + "!F{rr}", None),
        ("cwd", "closed_won_date (as exported)", "=IF(" + OP + "!G{rr}=\"\",\"\"," + OP + "!G{rr})", DATE),
        ("eff", "effective_start_date", "=IF(" + OP + "!H{rr}=\"\",\"\"," + OP + "!H{rr})", DATE),
        ("won", "Is a win? (stage governs)", "={stage}{r}=ST_WON", None),
        ("residue", "Residue: closed_won_date on a non-win (W2)", "=AND(NOT({won}{r}),{cwd}{r}<>\"\")", None),
        ("golive", "Account go-live (earliest Closed Won)", "=_xlfn.MINIFS(OP_EFF,OP_ACCT,{acct}{r},OP_STAGE,ST_WON)", DATE),
        ("latest", "Latest Closed Won effective date", "=_xlfn.MAXIFS(OP_EFF,OP_ACCT,{acct}{r},OP_STAGE,ST_WON)", DATE),
        ("isgl", "Is the go-live opportunity?", "=AND({won}{r},{eff}{r}={golive}{r})", None),
        ("key", "Lookup key", "={acct}{r}&\"|\"&IF({isgl}{r},1,0)", None),
        ("expan", "Later win dated in/after comp month on an older account", "=AND({won}{r},NOT({isgl}{r}),{eff}{r}>=CM_START,{golive}{r}<CM_START)", None),
    ]
    ot = Table(ws, 3, ocols)
    _write_offset(ot, N_OPPS, offset=-2)
    O = ot.L
    OR0, OR1 = 4, 3 + N_OPPS
    for key in ("key", "golive", "latest", "src", "residue", "expan", "acct", "won", "opp"):
        wb.defined_names["CO_" + key.upper()] = DefinedName("CO_" + key.upper(), attr_text=f"Calc_Opps!${O[key]}${OR0}:${O[key]}${OR1}")

    # ------------------------------------------------------------------ Calc_Accounts
    ws = S["Calc_Accounts"]
    ws.sheet_properties.tabColor = "548235"
    ws["A1"], ws["A1"].font = ("Calc_Accounts — one row per account (tab 7). New-logo windows (7.4–7.6), Existing revenue (7.7), "
                               "credits (7.2) and SPIFF book (7.12). AW = plan as written; REC = recommended (Treatments)."), F_SUB
    ws["A2"] = ("Shift counts come straight from tab 9 with COUNTIFS. Revenue per shift is an account property (7.3), "
                "so dollars = count x rate. Filter column 'Closer' or 'Owner' to trace any rep.")
    ws["A2"].font = F_NOTE
    MP = q("7. Opportunity to ID Mapping")
    W = "COUNTIFS(SH_ACCT,{app}{r},SH_WORKED,TRUE,SH_WDATE,\">=\"&{lo},SH_WDATE,\"<=\"&{hi})"
    Pall = "COUNTIFS(SH_ACCT,{app}{r},SH_PDATE,\">=\"&{lo},SH_PDATE,\"<=\"&{hi})"
    Pwk = "COUNTIFS(SH_ACCT,{app}{r},SH_WORKED,TRUE,SH_PDATE,\">=\"&{lo},SH_PDATE,\"<=\"&{hi})"
    LOW = "DATE(1900,1,1)"

    def cnt(kind, lo, hi):
        t = {"W": W, "P": Pall, "PW": Pwk}[kind]
        return t.replace("{lo}", lo).replace("{hi}", hi)

    def basis(lo, hi, rec):
        """Count of shifts for the account's credit basis between lo and hi (dates)."""
        gtm_all = cnt("P", lo, hi)
        gtm_wk = cnt("PW", lo, hi)
        nr = cnt("W", lo, hi)
        if rec:
            return f"IF({{gtm}}{{r}},IF({{hold}}{{r}},{gtm_wk},{gtm_all}),{nr})"
        return f"IF({{gtm}}{{r}},{gtm_all},{nr})"

    def win_cols(sfx, rec):
        lo1 = "IF({gtm}{r}," + LOW + ",{F}{r})"
        c = [
            ("n1" + sfx, f"W1 shifts ({sfx.upper().strip('_')})", "=IF(OR({c2}{r},{c3}{r})," + basis(lo1, "{e1}{r}", rec) + ",0)", INT),
            ("n2" + sfx, f"W2 shifts ({sfx.upper().strip('_')})", "=IF(OR({c2}{r},{c3}{r})," + basis("{e1}{r}+1", "{e2}{r}", rec) + ",0)", INT),
            ("n3" + sfx, f"W3 shifts ({sfx.upper().strip('_')})", "=IF({c3}{r}," + basis("{e2}{r}+1", "{e3}{r}", rec) + ",0)", INT),
            ("nj" + sfx, f"F1-30D shifts recognised in month ({sfx.upper().strip('_')})",
             "=IF({j1}{r}," + basis("MAX({F}{r},CM_START)", "MIN({e1}{r},CM_END)", rec).replace(
                 "IF({gtm}{r},", "IF({gtm}{r},").replace("\">=\"&MAX({F}{r},CM_START)", "\">=\"&MAX({F}{r},CM_START)") + ",0)", INT),
            ("nc" + sfx, f"GTM catch-up: W1 postings before month ({sfx.upper().strip('_')}) (A9)",
             "=IF(AND({gtm}{r},{fim}{r})," + (f"IF({{hold}}{{r}},{cnt('PW', LOW, 'CM_START-1')},{cnt('P', LOW, 'CM_START-1')})" if rec else cnt("P", LOW, "CM_START-1")) + ",0)", INT),
            ("njc" + sfx, f"F1-30D shifts on dates the closer was employed ({sfx.upper().strip('_')})",
             "=IF(AND({j1}{r},{cok}{r}),MAX(0," + basis("MAX({F}{r},CM_START,{ces}{r})", "MIN({e1}{r},CM_END,{cee}{r})", rec) + "),0)", INT),
        ]
        return c

    # For GTM, the F1-30D in-month postings run from the month start (not F): fix the lower bound
    def fix_gtm_lo(cols):
        out = []
        for k, h, t, f in cols:
            if k.startswith("nj") and not k.startswith("njc"):
                t = t.replace(cnt("P", "MAX({F}{r},CM_START)", "MIN({e1}{r},CM_END)"),
                              cnt("P", "CM_START", "MIN({e1}{r},CM_END)"))
                t = t.replace(cnt("PW", "MAX({F}{r},CM_START)", "MIN({e1}{r},CM_END)"),
                              cnt("PW", "CM_START", "MIN({e1}{r},CM_END)"))
            if k.startswith("njc"):
                t = t.replace(cnt("P", "MAX({F}{r},CM_START,{ces}{r})", "MIN({e1}{r},CM_END,{cee}{r})"),
                              cnt("P", "MAX(CM_START,{ces}{r})", "MIN({e1}{r},CM_END,{cee}{r})"))
                t = t.replace(cnt("PW", "MAX({F}{r},CM_START,{ces}{r})", "MIN({e1}{r},CM_END,{cee}{r})"),
                              cnt("PW", "MAX(CM_START,{ces}{r})", "MIN({e1}{r},CM_END,{cee}{r})"))
            out.append((k, h, t, f))
        return out

    def credit_cols(sfx):
        S_ = sfx
        return [
            ("w1t" + S_, f"F1-30D credit, total ({S_.strip('_').upper()})", "=({nj%s}{r}+{nc%s}{r})*{unit}{r}" % (S_, S_), USD),
            ("w1c" + S_, f"F1-30D credit to closer ({S_.strip('_').upper()})",
             "=({njc%s}{r}+IF(AND({fim}{r},{cokF}{r}),{nc%s}{r},0))*{unit}{r}" % (S_, S_), USD),
            ("w2t" + S_, f"F31-60D improvement credit ({S_.strip('_').upper()})",
             "=IF({c2}{r},MAX(0,{n2%s}{r}-{n1%s}{r})*{unit}{r}*LATER_FACTOR,0)" % (S_, S_), USD),
            ("w3t" + S_, f"F61-90D improvement credit ({S_.strip('_').upper()})",
             "=IF({c3}{r},MAX(0,{n3%s}{r}-MAX({n1%s}{r},{n2%s}{r}))*{unit}{r}*LATER_FACTOR,0)" % (S_, S_, S_), USD),
            ("newt" + S_, f"New credit total ({S_.strip('_').upper()})", "={w1t%s}{r}+{w2t%s}{r}+{w3t%s}{r}" % (S_, S_, S_), USD),
            ("newc" + S_, f"New credit to closer ({S_.strip('_').upper()})",
             "=IF({closer}{r}=\"\",0,{w1c%s}{r}+IF({cok2}{r},{w2t%s}{r},0)+IF({cok3}{r},{w3t%s}{r},0))" % (S_, S_, S_), USD),
            ("newu" + S_, f"New credit closer can't receive -> manager ({S_.strip('_').upper()})",
             "=IF({closer}{r}=\"\",0,{newt%s}{r}-{newc%s}{r})" % (S_, S_), USD),
            ("newx" + S_, f"New credit with blank closer -> nobody ({S_.strip('_').upper()})",
             "=IF({closer}{r}=\"\",{newt%s}{r},0)" % S_, USD),
        ]

    PPM = "MATCH({closer}{r},PP_CRM,0)"
    PPO = "MATCH({owner}{r},PP_CRM,0)"
    acols = [
        ("src", "Tab 7 row", "=ROW(" + MP + "!A{rr})", INT),
        ("acct", "crm_account_id", "=" + MP + "!A{rr}", None),
        ("app", "app_account_id", "=" + MP + "!B{rr}", None),
        ("bill", "billing_account_id", "=" + MP + "!C{rr}", None),
        ("orow", "Go-live opp: tab 8 row", "=INDEX(CO_SRC,MATCH({acct}{r}&\"|1\",CO_KEY,0))", INT),
        ("opp", "Go-live opportunity_id", "=INDEX(OP_ID,{orow}{r}-1)", None),
        ("nopp", "# opportunities", "=COUNTIF(OP_ACCT,{acct}{r})", INT),
        ("golive", "Go-live (7.3)", "=INDEX(CO_GOLIVE,MATCH({acct}{r},CO_ACCT,0))", DATE),
        ("latest", "Latest Closed Won effective (naive go-live)", "=INDEX(CO_LATEST,MATCH({acct}{r},CO_ACCT,0))", DATE),
        ("closer", "closed_by_owner_id", "=IF(INDEX(OP_CLOSER,{orow}{r}-1)=\"\",\"\",INDEX(OP_CLOSER,{orow}{r}-1))", None),
        ("cname", "Closer", "=IF({closer}{r}=\"\",\"(blank)\",IFERROR(INDEX(PP_NAME," + PPM + "),\"NOT ON ROSTER\"))", None),
        ("owner", "crm_owner_id", "=INDEX(OP_OWNER,{orow}{r}-1)", None),
        ("oname", "Owner", "=IFERROR(INDEX(PP_NAME," + PPO + "),\"NOT ON ROSTER\")", None),
        ("type", "account_type", "=INDEX(OP_TYPE,{orow}{r}-1)", None),
        ("gtm", "GTM? (A17)", "={type}{r}=GTM_TYPE", None),
        ("hold", "Treatment: GTM worked-only (A18)", "=COUNTIFS(TR_ACCT,{acct}{r},TR_CODE,\"GTM_WORKED_ONLY\",TR_ACTIVE,1)>0", None),
        ("seg", "segment (info only)", "=INDEX(OP_SEG,{orow}{r}-1)", None),
        ("charge", "Charge rate / shift", "=INDEX(OP_CHARGE,{orow}{r}-1)", USD0),
        ("pay", "Pay rate / shift", "=INDEX(OP_PAY,{orow}{r}-1)", USD0),
        ("margin", "Net Revenue / worked shift (7.1)", "={charge}{r}-{pay}{r}", USD0),
        ("terms", "Net terms", "=INDEX(OP_TERMS,{orow}{r}-1)", INT),
        ("tmult", "Terms multiplier (7.5)", "=INDEX(TERMS_MULT,MATCH({terms}{r},TERMS_DAYS,0))", PCT),
        ("crate", "Closer's Standard Payout", "=IFERROR(INDEX(PP_RATE," + PPM + "),0)", "0.00%"),
        ("unit", "Credit $ per basis shift", "=IF({gtm}{r},IF({crate}{r}=0,0,{charge}{r}*GTM_RATE/{crate}{r}),{margin}{r}*{tmult}{r})", USD,
         "Non-GTM: margin x terms multiplier. GTM: charge (GSV) x 3% / closer's rate, so it retires quota like any credit (7.6)."),
        ("ces", "Closer employed from", "=IFERROR(INDEX(PP_ES," + PPM + "),\"\")", DATE),
        ("cee", "Closer employed through", "=IFERROR(IF(INDEX(PP_EE," + PPM + ")=\"\",DATE(2999,12,31),INDEX(PP_EE," + PPM + ")),\"\")", DATE),
        ("ntot", "Shifts posted (all time in extract)", "=COUNTIF(SH_ACCT,{app}{r})", INT),
        ("F", "First worked shift (given, A15)", "=IFERROR(INDEX(SH_FIRST,MATCH({app}{r},SH_ACCT,0)),\"\")", DATE),
        ("e1", "W1 ends (day 30)", "=IF({F}{r}=\"\",\"\",{F}{r}+WIN_LEN-1)", DATE),
        ("e2", "W2 ends (day 60)", "=IF({F}{r}=\"\",\"\",{F}{r}+2*WIN_LEN-1)", DATE),
        ("e3", "W3 ends (day 90)", "=IF({F}{r}=\"\",\"\",{F}{r}+3*WIN_LEN-1)", DATE),
        ("j1", "W1 overlaps month?", "=IF({F}{r}=\"\",FALSE,AND({e1}{r}>=CM_START,IF({gtm}{r},TRUE,{F}{r}<=CM_END)))", None),
        ("fim", "First shift in month?", "=IF({F}{r}=\"\",FALSE,AND({F}{r}>=CM_START,{F}{r}<=CM_END))", None),
        ("c2", "W2 closes in month?", "=IF({F}{r}=\"\",FALSE,AND({e2}{r}>=CM_START,{e2}{r}<=CM_END))", None),
        ("c3", "W3 closes in month?", "=IF({F}{r}=\"\",FALSE,AND({e3}{r}>=CM_START,{e3}{r}<=CM_END))", None),
        ("cok", "Closer on roster?", "=AND({closer}{r}<>\"\",{ces}{r}<>\"\")", None),
        ("cokF", "Closer employed on first shift date?", "=IF(OR(NOT({cok}{r}),{F}{r}=\"\"),FALSE,AND({F}{r}>={ces}{r},{F}{r}<={cee}{r}))", None),
        ("cok2", "Closer employed when W2 closes?", "=IF(OR(NOT({cok}{r}),{F}{r}=\"\"),FALSE,AND({e2}{r}>={ces}{r},{e2}{r}<={cee}{r}))", None),
        ("cok3", "Closer employed when W3 closes?", "=IF(OR(NOT({cok}{r}),{F}{r}=\"\"),FALSE,AND({e3}{r}>={ces}{r},{e3}{r}<={cee}{r}))", None),
    ]
    aw_w = fix_gtm_lo(win_cols("_aw", False))
    rec_w = []
    for k, h, t, f in aw_w:
        body = t[1:].replace("COUNTIFS(SH_ACCT,{app}{r},SH_PDATE", "COUNTIFS(SH_ACCT,{app}{r},SH_WORKED,TRUE,SH_PDATE")
        rec_w.append((k.replace("_aw", "_rec"), h.replace("(AW)", "(REC)"),
                      "=IF(AND({gtm}{r},{hold}{r})," + body + ",{" + k + "}{r})", f))
    acols += aw_w + rec_w
    acols += credit_cols("_aw") + credit_cols("_rec")
    acols += [
        ("exist", "Existing in month? (go-live before month, 7.7)", "={golive}{r}<CM_START", None),
        ("oes", "Owner employed from", "=IFERROR(INDEX(PP_ES," + PPO + "),\"\")", DATE),
        ("oee", "Owner employed through", "=IFERROR(INDEX(PP_EEFF," + PPO + "),\"\")", DATE),
        ("xn", "Worked shifts in month", "=IF({F}{r}=\"\",0," + cnt("W", "CM_START", "CM_END") + ")", INT),
        ("xnc", "Worked shifts in month while owner employed", "=IF(AND({exist}{r},{oes}{r}<>\"\",{xn}{r}>0),MAX(0," + cnt("W", "MAX(CM_START,{oes}{r})", "{oee}{r}") + "),0)", INT),
        ("xnr", "Existing NR to owner", "=IF({exist}{r},{xnc}{r}*{margin}{r},0)", USD),
        ("xnru", "Existing NR owner can't receive -> manager", "=IF({exist}{r},({xn}{r}-{xnc}{r})*{margin}{r},0)", USD),
        ("cr", "Approved credits billed in month (7.2)", "=SUMIFS(BL_AMT,BL_ACCT,{bill}{r},BL_MONTH,CM_LABEL,BL_STATUS,ST_APPROVED)", USD),
        ("crd", "Denied credits billed in month (no effect)", "=SUMIFS(BL_AMT,BL_ACCT,{bill}{r},BL_MONTH,CM_LABEL,BL_STATUS,\"<>\"&ST_APPROVED)", USD),
        ("crx", "Credits applied to Existing", "=IF(AND({exist}{r},{oes}{r}<>\"\"),{cr}{r},0)", USD),
        ("crn", "Credits on a non-Existing account (A6: no effect)", "=IF({exist}{r},0,{cr}{r})", USD),
        ("xnet", "Existing NR net of credits to owner", "={xnr}{r}-{crx}{r}", USD),
        ("sbook", "In SPIFF book? (live before prior month)", "={golive}{r}<PM_START", None),
        ("sjune", "SPIFF: month ECNR net of credits", "=IF({sbook}{r},{xn}{r}*{margin}{r}-{cr}{r},0)", USD),
        ("smay", "SPIFF: prior month ECNR net of credits", "=IF({sbook}{r}," + cnt("W", "PM_START", "PM_END") + "*{margin}{r}-SUMIFS(BL_AMT,BL_ACCT,{bill}{r},BL_MONTH,PM_LABEL,BL_STATUS,ST_APPROVED),0)", USD),
        ("expan", "Ledger: later win in month on Existing account (E02)", "=AND({exist}{r},{latest}{r}>=CM_START)", None),
        ("expnr", "Ledger: Existing NR that naive go-live would drop", "=IF({expan}{r},{xn}{r}*{margin}{r}-{cr}{r},0)", USD),
    ]
    at = Table(ws, 4, acols)
    _write_offset(at, N_ACCTS, offset=-3)
    A = at.L
    AR0, AR1 = 5, 4 + N_ACCTS
    for k in list(A):
        wb.defined_names["CA_" + k.upper()] = DefinedName("CA_" + k.upper(), attr_text=f"Calc_Accounts!${A[k]}${AR0}:${A[k]}${AR1}")
    # header row lookups for the basis switch
    wb.defined_names["CA_HEAD"] = DefinedName("CA_HEAD", attr_text=f"Calc_Accounts!$A$4:${get_column_letter(len(acols))}$4")
    wb.defined_names["CA_ALL"] = DefinedName("CA_ALL", attr_text=f"Calc_Accounts!$A${AR0}:${get_column_letter(len(acols))}${AR1}")

    # ------------------------------------------------------------------ Pay sheets
    def pay_sheet(ws, basis_code, title):
        ws.sheet_properties.tabColor = "2F5597"
        ws["A1"], ws["A1"].font = title, F_SUB
        ws["A2"] = "Basis"
        ws["B2"] = basis_code
        ws["B2"].font = F_INPUT
        ws["C2"] = ("Both Pay_ sheets carry identical formulas; B2 picks the credit columns on Calc_Accounts "
                    "(AW = plan as written, REC = with Treatments).")
        ws["C2"].font = F_NOTE
        col = lambda base: f"INDEX(CA_ALL,0,MATCH(\"{base}\"&\"_\"&LOWER($B$2),CA_HEAD_KEYS,0))"
        prow = "{pr}"  # Calc_People row index (1-based)
        cols = [
            ("name", "Name", "=INDEX(PP_NAME,{k1})", None),
            ("role", "Role", "=INDEX(PP_ROLE,{k1})", None),
            ("hr", "hr_id", "=INDEX(PP_HR,{k1})", None),
            ("crm", "crm_owner_id", "=INDEX(PP_CRM,{k1})", None),
            ("l1", "L1", "=INDEX(PP_L1,{k1})", None),
            ("l2", "L2", "=INDEX(PP_L2,{k1})", None),
            ("l3", "L3", "=INDEX(PP_L3,{k1})", None),
            ("l4", "L4", "=INDEX(PP_L4,{k1})", None),
            ("isae", "AE?", "=INDEX(PP_ISAE,{k1})", None),
            ("isam", "AM?", "=INDEX(PP_ISAM,{k1})", None),
            ("newp", "New credit produced (7.4–7.6)", "=IF({isae}{r},SUMIFS(" + col("newc") + ",CA_CLOSER,{crm}{r}),0)", USD),
            ("newu", "New credit outside employment -> manager (7.11)", "=IF({isae}{r},SUMIFS(" + col("newu") + ",CA_CLOSER,{crm}{r}),0)", USD),
            ("exp", "Existing produced, net of credits (7.7)", "=IF({isam}{r},SUMIFS(CA_XNET,CA_OWNER,{crm}{r}),0)", USD),
            ("exu", "Existing outside employment -> manager", "=IF({isam}{r},SUMIFS(CA_XNRU,CA_OWNER,{crm}{r}),0)", USD),
            ("prod", "Produced", "=IF({isae}{r},{newp}{r},IF({isam}{r},{exp}{r},0))", USD),
            ("ramping", "Ramping?", "=INDEX(PP_RAMPING,{k1})", None),
            ("ramp", "Ramp floor", "=INDEX(PP_RAMPCR,{k1})", USD),
            ("icc", "IC credited (greater of, 7.10)", "=IF({ramping}{r},MAX({prod}{r},{ramp}{r}),{prod}{r})", USD),
            ("cn", "Roll-up contribution: New", "=IF({isae}{r},{icc}{r}+{newu}{r},0)", USD,
             "What managers carry for this rep: credited revenue after any ramp floor, plus credit the rep could not receive (7.8, 7.11)."),
            ("ce", "Roll-up contribution: Existing", "=IF({isam}{r},{icc}{r}+{exu}{r},0)", USD),
            ("carryn", "Carries: New (whole chain)", "=SUMIFS({cn}$5:{cn}$84,{l1}$5:{l1}$84,{hr}{r})+SUMIFS({cn}$5:{cn}$84,{l2}$5:{l2}$84,{hr}{r})+SUMIFS({cn}$5:{cn}$84,{l3}$5:{l3}$84,{hr}{r})+SUMIFS({cn}$5:{cn}$84,{l4}$5:{l4}$84,{hr}{r})", USD),
            ("carrye", "Carries: Existing (whole chain)", "=SUMIFS({ce}$5:{ce}$84,{l1}$5:{l1}$84,{hr}{r})+SUMIFS({ce}$5:{ce}$84,{l2}$5:{l2}$84,{hr}{r})+SUMIFS({ce}$5:{ce}$84,{l3}$5:{l3}$84,{hr}{r})+SUMIFS({ce}$5:{ce}$84,{l4}$5:{l4}$84,{hr}{r})", USD),
            ("c1", "Credited revenue (Metric I)", "=IF(OR({isae}{r},{isam}{r}),{icc}{r},IF(INDEX(PP_M1,{k1})=\"New Customer Net Revenue\",{carryn}{r},{carrye}{r}))", USD),
            ("c2", "Credited revenue (Metric II)", "=IF(INDEX(PP_M2,{k1})=\"Existing Customer Net Revenue\",{carrye}{r},0)", USD),
            ("q", "Quota prorated (Metric I)", "=INDEX(PP_QPR,{k1})", USD),
            ("att", "Attainment (Metric I)", "=IF({q}{r}>0,{c1}{r}/{q}{r},\"\")", PCT),
            ("t1", "Band I starts", "=INDEX(PP_T1,{k1})", USD0),
            ("t2", "Band II starts", "=INDEX(PP_T2,{k1})", USD0),
            ("acc", "Accelerators apply? (7.9)", "=AND(ISNUMBER({t1}{r}),NOT({ramping}{r}))", None),
            ("std", "$ in standard band", "=IF({acc}{r},MIN(MAX({c1}{r},0),{t1}{r}),{c1}{r})", USD),
            ("b1", "$ in band I", "=IF({acc}{r},MAX(0,MIN({c1}{r},{t2}{r})-{t1}{r}),0)", USD),
            ("b2", "$ in band II", "=IF({acc}{r},MAX(0,{c1}{r}-{t2}{r}),0)", USD),
            ("rate", "Standard Payout", "=INDEX(PP_RATE,{k1})", "0.0000%"),
            ("com1", "Commission Metric I", "=MAX(0,{rate}{r}*({std}{r}+IF({acc}{r},INDEX(PP_AM1,{k1})*{b1}{r}+INDEX(PP_AM2,{k1})*{b2}{r},0)))", USD,
             "Negative credited revenue pays $0 for the month (A7); the roll-up still uses the true figure."),
            ("com2", "Commission Metric II", "=INDEX(PP_RATE2,{k1})*{c2}{r}", USD),
            ("smay", "SPIFF book: prior month", "=IF({isam}{r},SUMIFS(CA_SMAY,CA_OWNER,{crm}{r}),0)", USD),
            ("sjun", "SPIFF book: comp month", "=IF({isam}{r},SUMIFS(CA_SJUNE,CA_OWNER,{crm}{r}),0)", USD),
            ("sratio", "SPIFF: month vs prior", "=IF({smay}{r}>0,{sjun}{r}/{smay}{r},\"\")", PCT),
            ("spiff", "Retention SPIFF (7.12)", "=IF(AND({isam}{r},NOT({ramping}{r}),{smay}{r}>0,{sjun}{r}>={smay}{r}),INDEX(PP_SPAMT,{k1}),0)", USD),
            ("usd", "Total payout USD", "=ROUND({com1}{r}+{com2}{r}+{spiff}{r},DEC)", USD),
            ("cur", "Currency", "=INDEX(PP_CUR,{k1})", None),
            ("fx", "FX (local per USD)", "=INDEX(PP_FX,{k1})", "0.0000"),
            ("loc", "Total payout local", "=ROUND({usd}{r}*{fx}{r},DEC)", NUM),
        ]
        t = Table(ws, 4, cols)
        _write_pay(t, N_REPS)
        return t

    # header-key row for basis lookups: a hidden row 3 on Calc_Accounts holding the column keys
    ws = S["Calc_Accounts"]
    for k, letter in A.items():
        ws[f"{letter}3"] = k
        ws[f"{letter}3"].font = Font(name=FONT, size=8, color="808080")
    wb.defined_names["CA_HEAD_KEYS"] = DefinedName("CA_HEAD_KEYS", attr_text=f"Calc_Accounts!$A$3:${get_column_letter(len(acols))}$3")

    PAW = pay_sheet(S["Pay_AsWritten"], "AW", "Pay_AsWritten — June payout applying Section 07 exactly as written (data errors handled by plan rules)")
    PRE = pay_sheet(S["Pay_Recommended"], "REC", "Pay_Recommended — as written plus the sign-off treatments on the Treatments tab")
    L = PAW.L
    for nm, sh in (("AW", "Pay_AsWritten"), ("REC", "Pay_Recommended")):
        for k in ("usd", "loc", "c1", "c2", "att", "icc", "prod", "carryn", "carrye", "spiff", "com1", "com2", "name",
                  "ramping", "role", "newu", "q", "sratio", "smay", "sjun", "cn", "ce", "b1", "b2"):
            wb.defined_names[f"{nm}_{k.upper()}"] = DefinedName(f"{nm}_{k.upper()}", attr_text=f"{sh}!${L[k]}$5:${L[k]}$84")

    _summary(S["Summary"], wb)
    _scenario(S["Scenario_Accelerator"])
    _ledger(S["Exceptions_Ledger"])
    _checks(S["Checks"])
    _observations(S["Plan_Observations"])
    _readme(S["README"])
    for wsn in wb.worksheets:
        wsn.sheet_view.zoomScale = 90
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print("saved", OUT)


def _write_people(t, n):
    t.ws.row_dimensions[3].height = 42
    first = t.hr + 1
    ws = t.ws
    for i, (key, header, tmpl, fmt, *rest) in enumerate(t.cols):
        c = ws.cell(row=t.hr, column=i + 1, value=header)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i + 1)].width = 14
    for k in range(n):
        r = first + k
        for i, (key, header, tmpl, fmt, *rest) in enumerate(t.cols):
            v = tmpl.format(r=r, rr=r - 2, **t.L)
            c = ws.cell(row=r, column=i + 1, value=v)
            c.font = F_BASE
            if fmt:
                c.number_format = fmt
    ws.column_dimensions["B"].width = 20
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(t.cols))}{first + n - 1}"


def _write_offset(t, n, offset):
    ws = t.ws
    ws.row_dimensions[t.hr].height = 54
    for i, (key, header, tmpl, fmt, *rest) in enumerate(t.cols):
        c = ws.cell(row=t.hr, column=i + 1, value=header)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if rest:
            c.comment = Comment(rest[0], "model")
        ws.column_dimensions[get_column_letter(i + 1)].width = 14
    for k in range(n):
        r = t.hr + 1 + k
        for i, (key, header, tmpl, fmt, *rest) in enumerate(t.cols):
            v = tmpl.format(r=r, rr=r + offset, **t.L)
            c = ws.cell(row=r, column=i + 1, value=v)
            c.font = F_BASE
            if fmt:
                c.number_format = fmt
    ws.freeze_panes = ws.cell(row=t.hr + 1, column=3)
    ws.auto_filter.ref = f"A{t.hr}:{get_column_letter(len(t.cols))}{t.hr + n}"


def _write_pay(t, n):
    ws = t.ws
    ws.row_dimensions[4].height = 54
    for i, (key, header, tmpl, fmt, *rest) in enumerate(t.cols):
        c = ws.cell(row=4, column=i + 1, value=header)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if rest:
            c.comment = Comment(rest[0], "model")
        ws.column_dimensions[get_column_letter(i + 1)].width = 13
    for k in range(n):
        r = 5 + k
        for i, (key, header, tmpl, fmt, *rest) in enumerate(t.cols):
            v = tmpl.format(r=r, k1=k + 1, **t.L)
            c = ws.cell(row=r, column=i + 1, value=v)
            c.font = F_BASE
            if fmt:
                c.number_format = fmt
    tot = 5 + n
    ws.cell(row=tot, column=1, value="TOTAL").font = F_BOLD
    for key in ("newp", "newu", "exp", "exu", "prod", "icc", "com1", "com2", "spiff", "usd"):
        letter = t.L[key]
        c = ws[f"{letter}{tot}"]
        c.value = f"=SUM({letter}5:{letter}{tot - 1})"
        c.font, c.number_format = F_BOLD, USD
    ws.column_dimensions["A"].width = 20
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(t.cols))}{tot - 1}"


# ---------------------------------------------------------------------- output tabs
def _hdr(ws, row, headers, widths=None):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=i + 1, value=h)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.row_dimensions[row].height = 42


def _summary(ws, wb):
    ws.sheet_properties.tabColor = "C00000"
    ws["A1"], ws["A1"].font = "June 2026 commissions — summary for all 80 reps", F_TITLE
    ws["A3"], ws["A3"].font = "Total payout, plan as written (USD)", F_BOLD
    ws["D3"] = "=SUM(AW_USD)"
    ws["A4"], ws["A4"].font = "Total payout, recommended for sign-off (USD)", F_BOLD
    ws["D4"] = "=SUM(REC_USD)"
    ws["A5"], ws["A5"].font = "Held pending investigation (as written − recommended)", F_BOLD
    ws["D5"] = "=D3-D4"
    ws["A6"], ws["A6"].font = "Model checks", F_BOLD
    ws["D6"] = "=Checks!C4"
    for c in ("D3", "D4", "D5"):
        ws[c].number_format, ws[c].font = USD, Font(name=FONT, size=12, bold=True)
    ws["D6"].font = Font(name=FONT, size=12, bold=True)
    ws["F3"] = "Recommended = as written, minus the GTM shift-template credit on one account (Exceptions E01), which is held — not forfeited — pending ops and leadership review."
    ws["F3"].font = F_NOTE
    ws["F4"] = "Local currency = USD x tab 3 rate (local units per 1 USD), rounded once per rep (A12)."
    ws["F4"].font = F_NOTE
    heads = ["Rep", "Role", "Manager", "Country", "Currency", "Quota (prorated)", "Credited revenue",
             "Attainment", "Ramping?", "Payout USD — as written", "Payout USD — recommended", "Difference USD",
             "Payout local — as written", "Payout local — recommended", "Exceptions / notes"]
    _hdr(ws, 8, heads, [20, 22, 18, 8, 9, 14, 15, 11, 9, 15, 15, 13, 16, 16, 60])
    notes = REP_NOTES
    for k in range(N_REPS):
        r = 9 + k
        i = k + 1
        f = {
            1: f"=INDEX(PP_NAME,{i})", 2: f"=INDEX(PP_ROLE,{i})", 3: f"=INDEX(PP_MGRNAME,{i})",
            4: f"=INDEX(PP_CC,{i})", 5: f"=INDEX(PP_CUR,{i})", 6: f"=INDEX(PP_QPR,{i})",
            7: f"=INDEX(REC_C1,{i})", 8: f"=INDEX(REC_ATT,{i})", 9: f"=IF(INDEX(REC_RAMPING,{i}),\"Yes\",\"\")",
            10: f"=INDEX(AW_USD,{i})", 11: f"=INDEX(REC_USD,{i})", 12: f"=J{r}-K{r}",
            13: f"=INDEX(AW_LOC,{i})", 14: f"=INDEX(REC_LOC,{i})",
        }
        for col, v in f.items():
            c = ws.cell(row=r, column=col, value=v)
            c.font = F_BASE
            c.number_format = {6: USD0, 7: USD0, 8: PCT, 10: USD, 11: USD, 12: USD, 13: NUM, 14: NUM}.get(col, "General")
        c = ws.cell(row=r, column=15, value=f'=IFERROR(INDEX(NOTE_TXT,MATCH(A{r},NOTE_NAME,0)),"")')
        c.font = F_NOTE
    tot = 9 + N_REPS
    ws.cell(row=tot, column=1, value="TOTAL").font = F_BOLD
    for col in ("J", "K", "L"):
        ws[f"{col}{tot}"] = f"=SUM({col}9:{col}{tot - 1})"
        ws[f"{col}{tot}"].font, ws[f"{col}{tot}"].number_format = F_BOLD, USD
    # notes table (documentation, typed) placed to the right
    ws["R8"], ws["S8"] = "Note: rep", "Note text"
    ws["R8"].font = ws["S8"].font = F_HEAD
    ws["R8"].fill = ws["S8"].fill = FILL_HEAD
    for j, (nm, txt) in enumerate(notes):
        ws.cell(row=9 + j, column=18, value=nm).font = F_INPUT
        ws.cell(row=9 + j, column=19, value=txt).font = F_INPUT
    wb.defined_names["NOTE_NAME"] = DefinedName("NOTE_NAME", attr_text=f"Summary!$R$9:$R${9 + len(notes) - 1}")
    wb.defined_names["NOTE_TXT"] = DefinedName("NOTE_TXT", attr_text=f"Summary!$S$9:$S${9 + len(notes) - 1}")
    ws.column_dimensions["R"].width = 18
    ws.column_dimensions["S"].width = 70
    ws.freeze_panes = "B9"
    ws.auto_filter.ref = f"A8:O{tot - 1}"


REP_NOTES = [
    ("Jon Snow", "E01: $219,100 GSV credited on a rural GTM account where 268 of 313 posted shifts were never worked (AE-loaded shift templates). Recommended credits worked shifts only; rest held."),
    ("Bruce Wayne", "E01 roll-up: carries Jon Snow's GTM credit."),
    ("Ebenezer Scrooge", "E01 roll-up via Bruce Wayne."),
    ("Hannibal Lecter", "E01 roll-up; Metric I New + Metric II Existing."),
    ("Katniss Everdeen", "E04: left 12 Jun. Credited only for credit recognised to 12 Jun; quota prorated 12/30. Rest carried by her manager."),
    ("Severus Snape", "E04: left 29 May, before the month. Earns nothing; Lord Voldemort carries his accounts' June credit."),
    ("Tom Sawyer", "E04: joined 15 Jun. Credit before 15 Jun goes to his manager. Quota prorated; bands are absolute (7.9). Ramping."),
    ("Ron Weasley", "E04: joined 30 Jun. One ramp day; earlier credit goes to his manager."),
    ("Lord Voldemort", "E07: HR role start 1 May vs previous role end 21 Jun (Slack: took over 22 Jun). Snapshot manager line held all month (A-2): carries the team for June. $0 impact."),
    ("Jay Gatsby", "E07: HR role start 1 May vs SMB AE role end 7 Jun (Slack: moved 8 Jun). Whole month at ENT (A11). Ramping either way; $0 impact."),
    ("Scout Finch", "E07: HR role start 1 May vs previous end 31 May. Ramping (promotion into ENT AE): no accelerators (7.9). Paid on produced, which exceeds the floor."),
    ("Walter White", "E09: HR start 2021 vs Slack 'started two weeks ago'. HR governs (A16); HR to confirm."),
    ("Frodo Baggins", "Ramping: 100% floor days 1-30, 50% days 31-60, so June splits on 17 Jun. Floor applied."),
    ("Gimli", "Rehired 4 May: ramp restarts (7.10). Produced exceeds the 50% floor."),
    ("Anna", "Ramping after move to ENT AM on 11 May: credited at the full-quota floor ($450k vs $315k produced); no SPIFF while ramping. Plan weakness P2."),
    ("Percy Jackson", "Ramping after AE->AM move on 1 May: floor exceeds book; no SPIFF while ramping. Plan weakness P2."),
    ("Dorothy Gale", "The GTM account in her book went live in June, so it is not in her retention book. SPIFF earned."),
    ("Dr. Watson", "SPIFF missed: June book at ~90% of May on the same accounts."),
]


def _scenario(ws):
    ws.sheet_properties.tabColor = "C00000"
    ws["A1"], ws["A1"].font = "Plan change: SMB AE accelerator 1.2x / 1.3x  ->  1.3x / 1.5x (task 3.3)", F_TITLE
    ws["A2"] = ("Repricing changes only the SMB AE accelerator multipliers (Control F24:G24). Credited revenue, ramp status and "
                "every other role are unchanged: managers carry revenue, not commission, so they do not move.")
    ws["A2"].font = F_NOTE
    ws["A4"], ws["A4"].font = "Answer", F_SUB
    lab = [("June as paid (corrected, recommended basis — A3)", "=Summary!D4"),
           ("June repriced", "=B5+B7"),
           ("Delta", "=SUM(L12:L91)"),
           ("Delta on the plan-as-written basis (incl. E01 account)", "=SUM(K12:K91)"),
           ("Reps whose pay changes", "=COUNTIF(L12:L91,\">0\")"),
           ("Share of delta from the single largest beneficiary", "=IF(B7=0,0,MAX(L12:L91)/B7)")]
    for i, (l, f) in enumerate(lab):
        ws.cell(row=5 + i, column=1, value=l).font = F_BASE
        c = ws.cell(row=5 + i, column=2, value=f)
        c.font = F_BOLD
        c.number_format = PCT if i == 5 else (INT if i == 4 else USD)
    ws["B7"].fill = FILL_KEY
    heads = ["Rep", "Role", "Ramping?", "Credited (REC)", "Credited (AW)", "$ in band I (REC)", "$ in band II (REC)",
             "Commission now (REC)", "Commission repriced (REC)", "Commission repriced (AW)", "Delta (AW)", "Delta (REC)"]
    _hdr(ws, 11, heads, [42, 22, 10, 14, 14, 14, 14, 15, 15, 15, 12, 12])
    for k in range(N_REPS):
        r = 12 + k
        i = k + 1
        def rep(c1, b1, b2):
            return (f"IF(AND(INDEX(PP_ROLE,{i})=\"SMB Account Executive\",NOT(INDEX(PP_RAMPING,{i}))),"
                    f"INDEX(PP_RATE,{i})*(MIN({c1},INDEX(PP_T1,{i}))+SCN_M1*MAX(0,MIN({c1},INDEX(PP_T2,{i}))-INDEX(PP_T1,{i}))"
                    f"+SCN_M2*MAX(0,{c1}-INDEX(PP_T2,{i}))),NA())")
        vals = {
            1: f"=INDEX(PP_NAME,{i})", 2: f"=INDEX(PP_ROLE,{i})", 3: f"=IF(INDEX(PP_RAMPING,{i}),\"Yes\",\"\")",
            4: f"=INDEX(REC_C1,{i})", 5: f"=INDEX(AW_C1,{i})",
            6: f"=INDEX(REC_B1,{i})", 7: f"=INDEX(REC_B2,{i})",
            8: f"=INDEX(REC_COM1,{i})",
            9: f"=IFERROR({rep(f'D{r}', '', '')},H{r})",
            10: f"=IFERROR({rep(f'E{r}', '', '')},INDEX(AW_COM1,{i}))",
            11: f"=J{r}-INDEX(AW_COM1,{i})", 12: f"=I{r}-H{r}",
        }
        for col, v in vals.items():
            c = ws.cell(row=r, column=col, value=v)
            c.font = F_BASE
            if col >= 4:
                c.number_format = USD
    ws.auto_filter.ref = "A11:L91"
    ws.freeze_panes = "B12"
    ws["D4"], ws["D4"].font = "Assumptions", F_SUB
    for j, t in enumerate([
        "1. Baseline is our corrected June on the recommended basis (A3), not the statements that went out.",
        "2. Same credited revenue: reps are assumed not to change behaviour in a month already worked.",
        "3. Ramping SMB AEs stay ineligible for accelerators (7.9); only non-ramping SMB AEs reprice.",
        "4. Bands stay absolute ($30k / $40k) and unprorated; only the multipliers change.",
        "5. Manager, director and Head pay is on revenue, so it does not change.",
    ]):
        ws.cell(row=5 + j, column=4, value=t).font = F_NOTE


def _ledger(ws):
    ws.sheet_properties.tabColor = "C00000"
    ws["A1"], ws["A1"].font = "Data quality exceptions ledger — ranked by payout impact (task 3.2)", F_TITLE
    ws["A2"] = ("Records / $ columns are live formulas on the model. 'Payout $ moved' compares the treated payout with a naive "
                "processing of the same rows (A2); where a naive run would need a second full model, the figure comes from the "
                "independent reference calculation (scripts/reference_calc.py) and is shown in blue with its method.")
    ws["A2"].font = F_NOTE
    heads = ["Rank", "ID", "Finding", "Root cause", "Records touched", "Credit $ moved", "Payout $ moved", "Reps affected",
             "Naive processing would…", "Treatment applied and why", "Category", "Decision"]
    _hdr(ws, 4, heads, [6, 6, 38, 38, 12, 14, 14, 26, 34, 48, 16, 14])
    rows = LEDGER_ROWS
    for i, row in enumerate(rows):
        r = 5 + i
        for j, v in enumerate(row):
            c = ws.cell(row=r, column=j + 1, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.font = F_BASE
            if j in (5, 6):
                c.number_format = USD
                if isinstance(v, (int, float)):
                    c.font = F_INPUT
            if j == 4:
                c.number_format = INT
    ws.freeze_panes = "C5"
    n = len(rows)
    ws.cell(row=6 + n, column=3, value="Largest finding vs all others combined (payout $)").font = F_BOLD
    ws.cell(row=6 + n, column=7, value=f"=G5-SUM(G6:G{4 + n})").number_format = USD
    ws.cell(row=6 + n, column=8, value="positive = the top finding outweighs the rest").font = F_NOTE


LEDGER_ROWS = [
    (1, "E01", "GTM shift-template inflation on account b2b39ae1 (closer Jon Snow, owner Dorothy Gale)",
     "GTM credits GSV on posted shifts (7.6). The AE loaded shift templates for a rural site with no scheduler (Slack 15 Jun); 268 of 313 postings never filled (14% fill vs ~78% norm). The facility's DON says they post a handful a day.",
     "=INDEX(CA_NTOT,MATCH(\"b2b39ae1-3d7a-47ee-962b-f400aabcd70c\",CA_ACCT,0))",
     "=SUM(CA_NEWT_AW)-SUM(CA_NEWT_REC)", "=Summary!D5",
     "Jon Snow, Bruce Wayne, Ebenezer Scrooge, Hannibal Lecter",
     "Pay ~$35k of credited revenue on shifts nobody worked, pushing Jon into accelerator band II.",
     "Recommended layer credits GSV of worked shifts only; unworked postings held pending ops/leadership review (A18). Plan-as-written figure kept visible. Fix: cap open postings, or pay GTM on filled shifts.",
     "Would not sign off", "A18 approved"),
    (2, "E02", "Expansion / renewal Closed Won dated in June on accounts live for months",
     "7.3: go-live is the earliest Closed Won; later wins do not restart it. Reading the latest (or any June) win as go-live makes these accounts look new.",
     "=COUNTIF(CA_EXPAN,TRUE)", "=SUM(CA_EXPNR)", 3320.95,
     "26 AMs and AM managers/directors (Lucy Pevensie +$369 because her SPIFF book changes)",
     "Drop the accounts from Existing in June: AMs and their chain lose the revenue; SPIFF books shift.",
     "Go-live = MINIFS of Closed Won effective date (Calc_Opps). Payout $: reference variant 'latest Closed Won as go-live'.",
     "Data / process", "Plan rule"),
    (3, "E04", "Credit recognised outside the closer's employment (leavers Severus 29 May, Katniss 12 Jun; joiners Tom 15 Jun, Ron 30 Jun)",
     "7.11: a rep is credited only where employed on the date the credit is recognised; the manager still carries it.",
     "=COUNTIF(CA_NEWU_REC,\">0\")", "=SUM(CA_NEWU_REC)", 1436.93,
     "Severus, Katniss (paid more naively); Captain Ahab, Lord Voldemort, Scrooge, Hannibal",
     "Pay the leaver/joiner for credit outside their dates.",
     "Credit split by shift date / window-close date vs HR employment; the un-creditable part rolls up to the manager. Payout $: reference variant 'ignore employment dates'.",
     "Plan rule", "Plan rule"),
    (4, "E10", "GTM postings made before the first worked shift, in an already-closed month",
     "7.6: windows have an upper bound only, so pre-first-shift postings fall in F1-30D, but their posting month closed before the window existed.",
     "=SUMPRODUCT((CA_FIM=TRUE)*(CA_GTM=TRUE)*CA_NC_REC)", "=SUMPRODUCT(CA_NC_REC,CA_UNIT)", 621.03,
     "Jon Snow, Daenerys Targaryen + roll-ups",
     "Drop them (May is closed) or restate May.",
     "Caught up in June, the first month the credit was knowable (A9). Payout $: reference variant 'no catch-up'.",
     "Plan gap", "A9 approved"),
    (5, "E06", "Denied billing tickets",
     "7.2: denied tickets change nothing.",
     "=COUNTIFS(BL_STATUS,\"<>\"&ST_APPROVED,BL_MONTH,CM_LABEL)", "=SUM(CA_CRD)", 533.27,
     "33 AMs and AM chain (Simba asked)",
     "Deduct every ticket, approved or not, from Existing revenue.",
     "Only status = approved reduces Existing revenue. Payout $: reference variant 'deduct denied'.",
     "Plan rule", "Plan rule"),
    (6, "E05", "Blank closed_by_owner_id on two June go-lives",
     "7.4: credit with no owner is not reassigned; 7.8: nobody carries it.",
     "=COUNTIF(CA_CNAME,\"(blank)\")", "=SUM(CA_NEWX_REC)", 0, "None (credit unassigned)",
     "Credit the AM owner or a manager by default.",
     "No one credited ($ shown is credit left unpaid). Sales ops to fix CRM ownership before July so later windows are paid correctly.",
     "Data error", "Plan rule"),
    (7, "E08", "Approved credits on accounts that are not Existing in the month",
     "7.2 + A6: credits reduce Existing only; new-logo credit is measured before credits.",
     "=COUNTIF(CA_CRN,\">0\")", "=SUM(CA_CRN)", 0, "None",
     "Reduce the AE's new-logo credit, or the AM anyway.",
     "No effect on pay; listed so finance sees credits not charged to anyone (A6).",
     "Plan rule", "A6 approved"),
    (8, "E03", "closed_won_date populated on Negotiation / Closed Lost rows",
     "Residue of deals moved to Closed Won and reversed (7.3).",
     "=COUNTIF(CO_RESIDUE,TRUE)", 0, 0, "None in June (Legolas asked)",
     "Treat as wins: could set go-live or create credit.",
     "opportunity_stage governs; residue ignored. Verified $0 effect in June.",
     "Data error", "Plan rule"),
    (9, "E07", "HR role start earlier than previous role end (Jay Gatsby, Scout Finch, Lord Voldemort)",
     "HR record not updated when the move happened; Slack gives 8 Jun (Jay) and 22 Jun (Voldemort).",
     "=COUNTIF(Calc_People!L4:L83,\"CHECK\")", 0, 0, "Jay, Scout, Voldemort (verified $0)",
     "Restart ramp / prorate on the wrong date.",
     "Current Role Start Date used (A-2, A19); HR to correct before July.",
     "Data error", "A19 approved"),
    (10, "E09", "Walter White: HR start 2021 vs Slack 'started two weeks ago'",
     "HR and reality may disagree (A-1).", 1, 0, 0, "Walter White",
     "Pay as tenured with no check.", "HR governs (A16, A20); ask HR to confirm before July. Not sized: no alternative start date on record.",
     "Data question", "A20 approved"),
]


def _checks(ws):
    ws.sheet_properties.tabColor = "70AD47"
    ws["A1"], ws["A1"].font = "Checks — every row must show PASS before sign-off", F_TITLE
    ws["A4"], ws["A4"].font = "Overall", F_BOLD
    ws["C4"] = '=IF(COUNTIF(D7:D40,"FAIL")=0,"ALL PASS","CHECK FAILURES: "&COUNTIF(D7:D40,"FAIL"))'
    ws["C4"].font = Font(name=FONT, size=12, bold=True)
    _hdr(ws, 6, ["#", "Check", "Difference / value", "Result", "What it proves"], [5, 60, 18, 10, 70])
    checks = [
        ("Reps on roster = 80 and each found once in HR", "=COUNTA('1. Sales Roster'!A2:A1000)-80+COUNTIF(Calc_People!F4:F83,\"MISSING\")", "Population is complete."),
        ("Accounts on ID map all found in tab 8", "=SUMPRODUCT(--(COUNTIF(OP_ACCT,CA_ACCT)=0))", "Every account has opportunities."),
        ("Every shift on tab 9 belongs to a mapped account", "=COUNTA(SH_ID)-SUM(CA_NTOT)", "No shift falls outside the model."),
        ("Every credit ticket belongs to a mapped account", "=COUNTA(BL_SHIFT)-SUMPRODUCT(COUNTIF(BL_ACCT,CA_BILL))", "No credit falls outside the model."),
        ("Worked shifts in month: accounts vs tab 9", "=SUM(CA_XN)-COUNTIFS(SH_WORKED,TRUE,SH_WDATE,\">=\"&CM_START,SH_WDATE,\"<=\"&CM_END)", "Monthly worked volume reconciles."),
        ("Approved credits in month: applied + not applied = tab 10", "=SUM(CA_CRX)+SUM(CA_CRN)-SUMIFS(BL_AMT,BL_MONTH,CM_LABEL,BL_STATUS,ST_APPROVED)", "Every approved credit is accounted for (A6)."),
        ("New credit (AW): closer + manager + nobody = total", "=SUM(CA_NEWC_AW)+SUM(CA_NEWU_AW)+SUM(CA_NEWX_AW)-SUM(CA_NEWT_AW)", "No new-logo dollar is lost or doubled."),
        ("New credit (REC): closer + manager + nobody = total", "=SUM(CA_NEWC_REC)+SUM(CA_NEWU_REC)+SUM(CA_NEWX_REC)-SUM(CA_NEWT_REC)", "Same for the recommended layer."),
        ("Head carries New = Director of Sales carries", "=SUMIFS(REC_C1,REC_ROLE,\"Head of Sales & AM\")-SUMIFS(REC_C1,REC_ROLE,\"Director of Sales\")", "Roll-up chain is complete for New."),
        ("Head carries Existing = both AM directors", "=SUMIFS(REC_C2,REC_ROLE,\"Head of Sales & AM\")-SUMIFS(REC_C1,REC_ROLE,\"Director of SMB AM\")-SUMIFS(REC_C1,REC_ROLE,\"Director of ENT AM\")", "Roll-up chain is complete for Existing."),
        ("Director of Sales = all AE contributions", "=SUMIFS(REC_C1,REC_ROLE,\"Director of Sales\")-SUM(REC_CN)", "Every AE dollar reaches the top of the New chain."),
        ("AM directors = all AM contributions", "=SUMIFS(REC_C1,REC_ROLE,\"Director of SMB AM\")+SUMIFS(REC_C1,REC_ROLE,\"Director of ENT AM\")-SUM(REC_CE)", "Every AM dollar reaches the top of the Existing chain."),
        ("Every rep has an FX rate", "=COUNTIF(PP_FX,\"NO RATE\")", "Local payout can be computed for all 80."),
        ("Summary totals = Pay sheets", "=Summary!J89-SUM(AW_USD)+Summary!K89-SUM(REC_USD)", "Output ties to calculation."),
        ("Recommended ≤ as written for every rep", "=SUMPRODUCT(--(REC_USD>AW_USD+0.005))", "Treatments only hold money; they never add it."),
        ("Independent recompute (Python) — recommended total", "=ROUND(SUM(REC_USD)-Checks!F22,2)", "Excel ties to scripts/reference_calc.py to the cent."),
        ("Independent recompute (Python) — as-written total", "=ROUND(SUM(AW_USD)-Checks!F23,2)", "Excel ties to scripts/reference_calc.py to the cent."),
    ]
    for i, (lab, f, why) in enumerate(checks):
        r = 7 + i
        ws.cell(row=r, column=1, value=i + 1).font = F_BASE
        ws.cell(row=r, column=2, value=lab).font = F_BASE
        c = ws.cell(row=r, column=3, value=f); c.font = F_BASE; c.number_format = NUM
        ws.cell(row=r, column=4, value=f'=IF(ABS(C{r})<0.005,"PASS","FAIL")').font = F_BOLD
        ws.cell(row=r, column=5, value=why).font = F_NOTE
    ws["E21"], ws["E21"].font = "Independent recompute totals (from scripts/reference_calc.py)", F_BOLD
    ws["E22"], ws["E23"] = "Recommended total (Python)", "As-written total (Python)"
    ws["F22"], ws["F23"] = REF_TOTALS["rec"], REF_TOTALS["aw"]
    for c in ("F22", "F23"):
        ws[c].font, ws[c].number_format = F_INPUT, USD
    ws["E24"] = "Source: python scripts/reference_calc.py (pandas re-implementation of Section 07, built separately from the formulas)."
    ws["E24"].font = F_NOTE
    ws.column_dimensions["F"].width = 16


REF_TOTALS = {"rec": 0.0, "aw": 0.0}


def _observations(ws):
    ws.sheet_properties.tabColor = "7030A0"
    ws["A1"], ws["A1"].font = "Plan observations — strengths, weaknesses and recommended changes", F_TITLE
    _hdr(ws, 3, ["#", "Observation", "Evidence in June", "Effect on behaviour / cost", "Recommendation"], [5, 40, 45, 45, 45])
    for i, row in enumerate(OBS):
        for j, v in enumerate(row):
            c = ws.cell(row=4 + i, column=j + 1, value=v)
            c.font = F_BASE
            c.alignment = Alignment(wrap_text=True, vertical="top")


OBS = [
    ("P3", "GTM pays on GSV of posted shifts, filled or not", "One account: 313 postings, 45 worked; $41k credited (E01)", "Rewards posting volume, which the AE can create; no cap on open shifts", "Pay GTM on filled shifts (or GSV x fill-rate floor); cap open postings; alert on fill < 50%"),
    ("P1", "AE windows (from first shift) overlap AM Existing credit (from go-live)", "Accounts live in May are Existing in June while the AE is still in F1-30D/F31-60D", "Same dollars pay twice in early months; unclear ownership", "Start Existing credit when the AE's 90 days end, or split explicitly"),
    ("P2", "Ramp floors roll up to managers", "Anna credited $450k vs $315k produced; flows to her manager, director and Head", "Managers paid on revenue nobody produced", "Roll up produced revenue, not the floor"),
    ("P4", "Improvement test on windows 2 and 3", "Credit only for growth over the best earlier window", "Rewards a soft first window (sandbagging)", "Credit windows 2-3 on absolute revenue at a lower rate"),
    ("P5", "Net 15 at 120%", "Net 15 accounts credit above their revenue (Luke Skywalker asked)", "Premium for terms the AE may concede in exchange for price", "Tie the multiplier to collected cash or cap it at 100%"),
    ("P6", "Retention SPIFF is all-or-nothing at 100%, and the AM manages credit timing", "Dr. Watson at ~90% gets $0; a small credit moved a month flips the result", "Cliff + control over the input = gaming risk", "Graduated SPIFF (e.g. 95-100-105%) and credits dated by shift month"),
    ("P7", "Window credits pay only when the window closes", "F31-60D/F61-90D land up to two months after the work", "Harder for reps to see cause and effect", "Show provisional window progress on each statement"),
    ("P8", "Strength: rules are explicit and mostly mechanical", "Every June payout is reproducible from source rows", "Low dispute risk once communicated", "Publish a one-page rules digest with each statement"),
]


def _readme(ws):
    ws.sheet_properties.tabColor = "000000"
    lines = [
        ("June 2026 Commission Model", F_TITLE),
        ("Healthcare staffing marketplace — Sales Compensation Analyst case. Built to be rerun by someone who is not the author.", F_NOTE),
        ("", None),
        ("Headline", F_SUB),
        ("=\"Total June payout, recommended: \"&TEXT(Summary!D4,\"$#,##0.00\")&\"   |   plan as written: \"&TEXT(Summary!D3,\"$#,##0.00\")&\"   |   held: \"&TEXT(Summary!D5,\"$#,##0.00\")", F_BOLD),
        ("=\"Largest finding: E01 GTM shift-template inflation (\"&TEXT(Exceptions_Ledger!G5,\"$#,##0\")&\" of payout). Model checks: \"&Checks!C4", F_BASE),
        ("", None),
        ("Tab map (left to right)", F_SUB),
        ("Outputs (red): Summary · Exceptions_Ledger · Scenario_Accelerator · Plan_Observations", F_BASE),
        ("Checks (green): Checks — reconciliations; all must PASS", F_BASE),
        ("Inputs you may edit (orange): Control (comp month, plan parameters, scenario lever) · Treatments (sign-off holds)", F_BASE),
        ("Calculations (dark green / blue): Calc_People → Calc_Opps → Calc_Accounts → Pay_AsWritten / Pay_Recommended", F_BASE),
        ("Source data (brown): tabs 1–11, exactly as received. Never edited; formulas read them through named ranges.", F_BASE),
        ("", None),
        ("How a payout is built", F_SUB),
        ("1. Calc_People: HR role, dates, manager chain, quota, rates, ramp floor and bands for each of the 80 reps.", F_BASE),
        ("2. Calc_Opps: go-live = earliest Closed Won effective date per account; stage governs, residue dates ignored.", F_BASE),
        ("3. Calc_Accounts: per account, shift counts from tab 9 (COUNTIFS) x revenue per shift → F1-30D / F31-60D / F61-90D credit (net terms or GTM), Existing revenue, credits, SPIFF book.", F_BASE),
        ("4. Pay_ sheets: sum each rep's accounts, apply ramp floor, roll up the chain, apply bands and accelerators, SPIFF, round, convert.", F_BASE),
        ("", None),
        ("How to trace any payout to source rows", F_SUB),
        ("a) Find the rep on Pay_Recommended (column 'Credited revenue').", F_BASE),
        ("b) On Calc_Accounts, filter 'Closer' (AE) or 'Owner' (AM) to the rep: each row shows the go-live opportunity's tab 8 row, the tab 7 row, the window dates and shift counts.", F_BASE),
        ("c) On tab 9, filter app_account_id to that account and the window dates: the rows counted are exactly those shifts. Credits: filter tab 10 by billing_account_id.", F_BASE),
        ("", None),
        ("Run next month (July)", F_SUB),
        ("1. Paste the new extracts into tabs 1–11 (same columns, header in row 1). Ranges accept up to 250,000 shifts.", F_BASE),
        ("2. Control!B5: set the comp month to 2026-07-01. Nothing else is date-specific.", F_BASE),
        ("3. Review Treatments (switch holds on/off) and add any new rep notes on Summary (columns R:S).", F_BASE),
        ("4. Recalculate; confirm Checks shows ALL PASS (update the two Python totals after rerunning scripts/reference_calc.py).", F_BASE),
        ("5. If tabs 1, 7 or 8 grow beyond 80 reps / 761 accounts / 815 opportunities, extend the Calc_ tables' rows (copy the last row down).", F_BASE),
        ("", None),
        ("Colour legend", F_SUB),
        ("Blue text = typed input or documented assumption · Black = formula · Yellow fill = key lever · Brown tabs = source data", F_BASE),
        ("", None),
        ("Key decisions (full log in PROJECT_PLAN.md)", F_SUB),
        ("A1 plan-as-written and recommended shown side by side · A6 credits on non-Existing accounts change nothing · A7 negative months pay $0 · A9 GTM pre-shift postings caught up in the first knowable month · A10 ramping = any ramp day in month · A11 month-end role for whole month · A12 round once · A15 first worked shift as given · A16 HR governs · A17 account_type GTM = tag · A18 GTM worked-only hold · A19 Current Role Start Date · A20 HR start governs", F_BASE),
    ]
    for i, (t, f) in enumerate(lines):
        c = ws.cell(row=1 + i, column=1, value=t)
        if f:
            c.font = f
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 150


if __name__ == "__main__":
    import json, sys
    ref = Path(__file__).resolve().parent / "ref_totals.json"
    if ref.exists():
        REF_TOTALS.update(json.loads(ref.read_text()))
    build()
